#!/usr/bin/env python3
"""
Sync the website's GUEST_LIST from the Google Sheet.

Usage:
    python3 _tools/sync_guests.py <sheet_export.xlsx> <index.html>

Source of truth: the 'Wedding Guest List' tab of "CJ & Sheilla-Wedding Prep".
Columns are found by their header text in row 5 (so moving a column is safe):
  - ATTENDEES          -> guest name   (rows 6 down)
  - RESERVED SEAT      -> seats        (blank -> 1)
  - Afterparty Invite  -> "Yes" means the guest is asked about the Alabang
                          afterparty on the site (written as ap:1)
A suffix like " + baby" is dropped from the name the site shows/submits
(the sheet's ATTENDING? formula ignores that suffix too, so RSVPs still match).

Nickname keys already on the site for a guest are kept; new guests get
their full name plus their first name (if no one else shares it).

Prints a summary and exits 0. First line is "CHANGED" or "NO CHANGES".
"""
import re, sys
import openpyxl

TAB, HEADER_ROW, FIRST_ROW = 'Wedding Guest List', 5, 6
DEFAULT_NAME_COL, DEFAULT_SEAT_COL = 3, 7          # C, G
START = 'var GUEST_LIST = ['
COMMENT = ("// GUEST LIST - auto-synced from the 'Wedding Guest List' tab of the "
           "Google Sheet. Edit the sheet, not this list. ap:1 = invited to the afterparty.")
ENTRY_RE = re.compile(r"\{name:'((?:[^'\\]|\\.)*)',seats:(\d+),keys:\[([^\]]*)\](,ap:1)?\}")


def norm(s):
    return re.sub(r'\s+', ' ', str(s).strip()).lower()


def display_name(raw):
    return re.sub(r'\s+', ' ', raw.split('+')[0]).strip()


def js_str(s):
    return "'" + s.replace('\\', '\\\\').replace("'", "\\'") + "'"


def find_columns(ws):
    name_col, seat_col, ap_col = DEFAULT_NAME_COL, DEFAULT_SEAT_COL, None
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(HEADER_ROW, c).value or '')
        if h == 'attendees':
            name_col = c
        elif 'reserved seat' in h:
            seat_col = c
        elif 'afterparty' in h and 'invite' in h:
            ap_col = c
    return name_col, seat_col, ap_col


def read_sheet(path):
    ws = openpyxl.load_workbook(path, data_only=True)[TAB]
    name_col, seat_col, ap_col = find_columns(ws)
    guests, seen = [], set()
    for r in range(FIRST_ROW, ws.max_row + 1):
        raw = ws.cell(r, name_col).value
        if raw is None or not str(raw).strip():
            continue
        name = display_name(str(raw))
        if not name or norm(name) in seen:
            continue
        seat = ws.cell(r, seat_col).value
        try:
            seats = int(float(seat)) if seat not in (None, '') else 1
        except (TypeError, ValueError):
            seats = 1
        ap = bool(ap_col) and norm(ws.cell(r, ap_col).value or '') == 'yes'
        guests.append((name, max(seats, 1), ap))
        seen.add(norm(name))
    return guests, ap_col is not None


def read_site(html):
    i = html.index(START)
    j = html.index('];', i)
    old = {}
    for m in ENTRY_RE.finditer(html[i:j]):
        keys = re.findall(r"'((?:[^'\\]|\\.)*)'", m.group(3))
        old[norm(m.group(1))] = (m.group(1), int(m.group(2)), keys, bool(m.group(4)))
    return i, j, old


def build_keys(name, existing, first_counts):
    keys = list(existing) if existing else []
    full = norm(name)
    if full not in keys:
        keys.insert(0, full)
    words = full.split(' ')
    if words[0].endswith('.') and len(words) > 1:          # "Ma. Josephine Carpio"
        alt = ' '.join(words[1:])
        if alt not in keys:
            keys.append(alt)
        words = words[1:]
    first = words[0]
    if len(words) > 1 and first_counts.get(first, 0) == 1 and first not in keys and not existing:
        keys.append(first)
    return keys


def main(xlsx, html_path):
    guests, has_ap_col = read_sheet(xlsx)
    html = open(html_path, encoding='utf-8').read()
    i, j, old = read_site(html)

    first_counts = {}
    for n, _, _ in guests:
        w = norm(n).split(' ')
        f = w[1] if w[0].endswith('.') and len(w) > 1 else w[0]
        first_counts[f] = first_counts.get(f, 0) + 1

    lines, added, seat_changes, ap_added, ap_removed = [], [], [], [], []
    for name, seats, ap in guests:
        prev = old.get(norm(name))
        keys = build_keys(name, prev[2] if prev else None, first_counts)
        if not prev:
            added.append(name)
        else:
            if prev[1] != seats:
                seat_changes.append(f'{name}: {prev[1]} -> {seats}')
            if ap and not prev[3]:
                ap_added.append(name)
            if prev[3] and not ap:
                ap_removed.append(name)
        lines.append('  {name:%s,seats:%d,keys:[%s]%s}' % (
            js_str(name), seats, ','.join(js_str(k) for k in keys), ',ap:1' if ap else ''))
    current = {norm(n) for n, _, _ in guests}
    removed = [v[0] for k, v in old.items() if k not in current]

    # Replace the comment line directly above the list (if it is a comment) and the list body
    line_start = html.rfind('\n', 0, i) + 1
    prev_line_start = html.rfind('\n', 0, line_start - 1) + 1
    prev_line = html[prev_line_start:line_start - 1]
    head_start = prev_line_start if prev_line.strip().startswith('//') else line_start
    new_block = COMMENT + '\n' + START + '\n' + ',\n'.join(lines) + '\n'
    new_html = html[:head_start] + new_block + html[j:]

    changed = new_html != html
    print('CHANGED' if changed else 'NO CHANGES')
    print(f'Guests on site: {len(guests)}  (total seats {sum(s for _, s, _ in guests)}; '
          f'afterparty invites {sum(1 for g in guests if g[2])})')
    if not has_ap_col:
        print("Warning: no 'Afterparty Invite' column found in row 5; nobody will be asked about the afterparty")
    for label, items in (('Added', added), ('Removed', removed), ('Seat changes', seat_changes),
                         ('Afterparty invite added', ap_added), ('Afterparty invite removed', ap_removed)):
        if items:
            print(f'{label}: ' + '; '.join(items))
    if changed:
        open(html_path, 'w', encoding='utf-8').write(new_html)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
